"""Multi-page documents: Word page splitting, page numbers in sources, duplicate uploads,
and retrieval that is not crowded out by identical chunks."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as new_docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.config import settings
from app.models.schemas import Query, RecognizedChunk
from app.modules import chunker, embedder, llm_answer, retrieval, text_parser, vector_store

PAGES = [
    "Unit 1. The OSI model has seven layers and TCP provides reliable delivery.",
    "Unit 2. Second normal form removes partial dependencies from a relation.",
    "Unit 3. Lexical analysis converts source characters into tokens for the parser.",
]


@pytest.fixture
def fake_llm(monkeypatch):
    fake = SimpleNamespace(reply="An answer.")
    monkeypatch.setattr(llm_answer, "_generate", lambda prompt: fake.reply)
    return fake


def _upload(client, path: Path, **form):
    with path.open("rb") as f:
        return client.post("/upload", files={"file": (path.name, f, "application/octet-stream")}, data=form)


def _write_paged_docx(path: Path, double_mark: bool = False) -> Path:
    """A Word file with a manual page break between PAGES. With double_mark, the paragraph after
    each break also carries Word's own layout mark, as a file saved by Word does."""
    doc = new_docx()
    for i, text in enumerate(PAGES):
        if i:
            doc.add_page_break()
        para = doc.add_paragraph(text)
        if i and double_mark:
            para.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    doc.save(path)
    return path


# --- Word pages -----------------------------------------------------------------------------


def test_docx_splits_at_page_breaks(tmp_path) -> None:
    pages = text_parser.parse(str(_write_paged_docx(tmp_path / "notes.docx")))
    assert [text for text, _ in pages] == PAGES


def test_docx_break_marked_twice_is_one_page_end(tmp_path) -> None:
    pages = text_parser.parse(str(_write_paged_docx(tmp_path / "notes.docx", double_mark=True)))
    assert len(pages) == 3


def test_docx_word_layout_marks_split_pages(tmp_path) -> None:
    doc = new_docx()
    doc.add_paragraph(PAGES[0])
    para = doc.add_paragraph(PAGES[1])
    para.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    doc.save(tmp_path / "saved_by_word.docx")
    assert len(text_parser.parse(str(tmp_path / "saved_by_word.docx"))) == 2


def test_docx_upload_reports_its_pages(client, tmp_path) -> None:
    resp = _upload(client, _write_paged_docx(tmp_path / "notes.docx"))
    assert resp.status_code == 201
    assert resp.json()["page_count"] == 3


# --- page numbers in sources ----------------------------------------------------------------


def test_chunk_pages_numbered_gives_one_based_pages() -> None:
    numbered = chunker.chunk_pages_numbered("d", [("first page", 1.0), ("", 1.0), ("third page", 1.0)])
    assert [(loc.page, c.chunk_id) for loc, c in numbered] == [(1, "d:0"), (3, "d:1")]


def test_sources_cite_the_page_the_answer_came_from(client, pdf_file, fake_llm) -> None:
    _upload(client, pdf_file(PAGES))
    body = {"query_id": "q-page", "question_text": "What does lexical analysis convert into tokens?", "user_id": "u"}
    answer = client.post("/ask", json=body).json()

    sources = client.get(f"/answer/{answer['answer_id']}/sources").json()

    assert sources[0]["page"] == 3
    assert "Lexical analysis" in sources[0]["text"]
    assert {s["page"] for s in sources} == {1, 2, 3}


# --- duplicate uploads ----------------------------------------------------------------------


def test_same_file_twice_is_stored_once(client, pdf_file) -> None:
    path = pdf_file(PAGES)
    first_resp, second_resp = _upload(client, path), _upload(client, path)
    first, second = first_resp.json(), second_resp.json()

    assert (first_resp.status_code, second_resp.status_code) == (201, 200)  # 200 = already there
    assert second["doc_id"] == first["doc_id"]
    assert len(client.get("/documents").json()) == 1
    assert len(list(settings.upload_dir.iterdir())) == 1


def test_same_file_with_a_different_format_hint_is_read_again(client, pdf_file, fake_htr) -> None:
    fake_htr.text, fake_htr.conf = "handwritten reading of the page", 0.9
    path = pdf_file(PAGES)
    first = _upload(client, path).json()
    second = _upload(client, path, format_hint="handwritten").json()

    assert second["doc_id"] != first["doc_id"]
    assert second["format_type"] == "handwritten"


def test_different_files_are_both_stored(client, pdf_file) -> None:
    _upload(client, pdf_file(PAGES, name="a.pdf"))
    _upload(client, pdf_file(PAGES[:2], name="b.pdf"))
    assert len(client.get("/documents").json()) == 2


# --- delete ---------------------------------------------------------------------------------


def test_delete_removes_record_file_and_chunks(client, pdf_file, fake_llm) -> None:
    doc_id = _upload(client, pdf_file(PAGES)).json()["doc_id"]
    [stored] = list(settings.upload_dir.iterdir())

    assert client.delete(f"/documents/{doc_id}").status_code == 204

    assert client.get("/documents").json() == []
    assert not stored.exists()
    assert vector_store.get_chunks([f"{doc_id}:0", f"{doc_id}:1"]) == {}
    answer = client.post("/ask", json={"query_id": "after", "question_text": "What is lexical analysis?", "user_id": "u"}).json()
    assert answer["answer_text"] == llm_answer.NO_DOCUMENTS_ANSWER


def test_delete_keeps_other_documents(client, pdf_file) -> None:
    keep = _upload(client, pdf_file(PAGES, name="keep.pdf")).json()["doc_id"]
    drop = _upload(client, pdf_file(PAGES[:1], name="drop.pdf")).json()["doc_id"]
    client.delete(f"/documents/{drop}")
    assert [d["doc_id"] for d in client.get("/documents").json()] == [keep]
    assert vector_store.get_chunks([f"{keep}:0"])


def test_delete_unknown_document_is_404(client) -> None:
    assert client.delete("/documents/nope").status_code == 404


def test_file_can_be_uploaded_again_after_delete(client, pdf_file) -> None:
    path = pdf_file(PAGES)
    first = _upload(client, path).json()["doc_id"]
    client.delete(f"/documents/{first}")
    again = _upload(client, path)
    assert again.status_code == 201 and again.json()["doc_id"] != first


def test_frontend_has_delete_button(client) -> None:
    html = client.get("/").text
    assert 'method: "DELETE"' in html and "Already uploaded" in html


# --- retrieval skips identical chunks -------------------------------------------------------


def _store(doc_id: str, texts: list[str]) -> None:
    chunks = [RecognizedChunk(chunk_id=f"{doc_id}:{i}", text=t, raw_conf=1.0, calibrated_conf=1.0) for i, t in enumerate(texts)]
    vector_store.add_chunks(doc_id, embedder.embed_chunks(chunks))


def test_retrieval_is_not_crowded_out_by_duplicate_documents(fake_embeddings, chroma, monkeypatch) -> None:
    monkeypatch.setattr(settings, "candidate_multiplier", 1)
    texts = [f"Topic {n}: notes about scheduling and memory, part {n}." for n in range(6)]
    for copy in range(5):  # the same document already stored five times
        _store(f"copy{copy}", texts)

    hits = retrieval.retrieve(Query(query_id="q", question_text="scheduling and memory notes", user_id="u"), top_k=5)

    assert len({chunk.text for chunk, _ in hits}) == 5


def test_identical_passages_keep_the_legible_copy(fake_embeddings, chroma) -> None:
    """Dropping repeats happens after the confidence-aware ranking, not before it."""
    slightly_closer = embedder.embed_texts(["exam dates march april"])[0]
    vector_store.add_chunks("doc", [
        RecognizedChunk(chunk_id="doc:0", text="exam dates march", raw_conf=0.3, calibrated_conf=0.3, vector=slightly_closer),
        RecognizedChunk(chunk_id="doc:1", text="exam dates march", raw_conf=0.95, calibrated_conf=0.95,
                        vector=embedder.embed_texts(["exam dates march"])[0]),
    ])
    hits = retrieval.retrieve(Query(query_id="q", question_text="exam dates march april", user_id="u"), top_k=2)
    assert [chunk.chunk_id for chunk, _ in hits] == ["doc:1"]


def test_retrieval_with_fewer_distinct_chunks_than_top_k(fake_embeddings, chroma) -> None:
    for copy in range(3):
        _store(f"copy{copy}", ["only one passage about paging"])
    hits = retrieval.retrieve(Query(query_id="q", question_text="paging", user_id="u"), top_k=5)
    assert len(hits) == 1


# --- line references ------------------------------------------------------------------------


def test_chunk_line_ranges_skip_blank_lines() -> None:
    page = "Heading\n\n  Indented second line\nThird line"
    [(loc, chunk)] = chunker.chunk_pages_numbered("d", [(page, 1.0)])
    assert (loc.page, loc.first_line, loc.last_line) == (1, 1, 3)


def test_long_page_chunks_cover_increasing_line_ranges() -> None:
    # Every word is unique ("L7w3" = line 7, word 3), so each chunk's true first/last line is known.
    page = "\n".join(" ".join(f"L{n}w{k}" for k in range(12)) for n in range(1, 31))
    numbered = chunker.chunk_pages_numbered("d", [(page, 1.0)], chunk_size=150, overlap=40)
    assert len(numbered) > 5
    for loc, chunk in numbered:
        words = chunk.text.split()
        assert loc.first_line == int(words[0][1:].split("w")[0])
        assert loc.last_line == int(words[-1][1:].split("w")[0])


def test_sources_cite_page_and_lines(client, pdf_file, fake_llm) -> None:
    _upload(client, pdf_file(["Intro line\nLexical analysis converts characters into tokens.\nMore text"]))
    body = {"query_id": "q-lines", "question_text": "What does lexical analysis convert?", "user_id": "u"}
    answer = client.post("/ask", json=body).json()
    [source] = client.get(f"/answer/{answer['answer_id']}/sources").json()
    assert (source["page"], source["first_line"], source["last_line"]) == (1, 1, 3)
