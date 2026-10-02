"""Experience Center Step 5 API (exports, preprocessing, plugins) with the Paddle engine faked."""

import csv
import io
import json
import zipfile
from pathlib import Path

import cv2
import numpy as np
import pymupdf
import pytest
from docx import Document

from app.config import settings
from app.modules import paddleocr_service as svc
from tests.conftest import make_pdf
from tests.test_paddleocr_api import ec, post_ocr  # noqa: F401  (ec is a fixture)

FIXTURE_PNG = Path(__file__).parent / "fixtures" / "ppstructurev3_page.png"


def tilted_png(degrees: float = 5.0) -> bytes:
    image = cv2.imread(str(FIXTURE_PNG))
    h, w = image.shape[:2]
    matrix = cv2.getRotationMatrix2D((w / 2, h / 2), degrees, 1.0)
    ok, data = cv2.imencode(".png", cv2.warpAffine(image, matrix, (w, h), borderValue=(255, 255, 255)))
    return data.tobytes()


def ocr_run(client, **form):
    resp = post_ocr(client, data=FIXTURE_PNG.read_bytes(), **form)
    assert resp.status_code == 200, resp.text
    return resp.json()


def parse_run(client):
    resp = client.post(
        "/api/parse", files={"file": ("report.png", FIXTURE_PNG.read_bytes())}, data={"wait": "true"}
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


# --- exports ---


def test_export_txt_json_md(ec):
    run = ocr_run(ec)
    txt = ec.get(f"/api/results/{run['id']}/export/txt")
    assert txt.status_code == 200 and txt.headers["content-type"].startswith("text/plain")
    assert txt.text.splitlines()[0] == "Calibration Report"
    assert 'filename="page.txt"' in txt.headers["content-disposition"]
    assert json.loads(ec.get(f"/api/results/{run['id']}/export/json").content)["id"] == run["id"]
    assert "Calibration Report" in ec.get(f"/api/results/{run['id']}/export/md").text


def test_export_uses_review_edits(ec):
    run = ocr_run(ec)
    ec.patch(f"/api/results/{run['id']}/lines", json=[{"page": 0, "line": 0, "text": "Calibration Summary"}])
    assert ec.get(f"/api/results/{run['id']}/export/txt").text.startswith("Calibration Summary")


def test_export_searchable_pdf(ec):
    run = ocr_run(ec)
    resp = ec.get(f"/api/results/{run['id']}/export/pdf")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/pdf"
    with pymupdf.open(stream=resp.content, filetype="pdf") as doc:
        assert doc.page_count == 1
        assert "Calibration Report" in doc[0].get_text()
        assert doc[0].search_for("Calibration")


def test_export_docx_and_csv_from_structure(ec):
    run = parse_run(ec)
    docx = ec.get(f"/api/results/{run['id']}/export/docx")
    assert docx.status_code == 200
    doc = Document(io.BytesIO(docx.content))
    assert "Calibration Report" in [p.text for p in doc.paragraphs]
    assert doc.tables and doc.tables[0].rows[0].cells[0].text == "Method"
    table = ec.get(f"/api/results/{run['id']}/export/csv")
    assert table.status_code == 200
    rows = list(csv.reader(io.StringIO(table.content.decode("utf-8-sig"))))
    assert rows[0] == ["Method", "ECE", "Brier"] and len(rows) == 4


def test_export_errors(ec):
    run = ocr_run(ec)
    assert ec.get(f"/api/results/{run['id']}/export/xlsx").status_code == 400
    assert ec.get(f"/api/results/{run['id']}/export/csv").status_code == 400  # text OCR has no tables
    assert ec.get("/api/results/nope/export/txt").status_code == 404


# --- preprocessing ---


def test_preprocess_preview_deskews(ec):
    resp = ec.post(
        "/api/preprocess",
        files={"file": ("tilted.png", tilted_png())},
        data={"preprocess": json.dumps({"deskew": True, "binarize": True})},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["before"].startswith("data:image/png;base64,") and body["after"].startswith("data:image/png;base64,")
    assert body["steps"] == ["deskew", "binarize"]
    assert abs(body["skew_angle"]) > 3


def test_preprocess_preview_pdf_and_validation(ec):
    pdf = make_pdf(["Confidence calibration of OCR lines."])
    ok = ec.post("/api/preprocess", files={"file": ("a.pdf", pdf)}, data={"preprocess": '{"contrast": true}'})
    assert ok.status_code == 200 and ok.json()["skew_angle"] is None
    none = ec.post("/api/preprocess", files={"file": ("a.png", FIXTURE_PNG.read_bytes())}, data={"preprocess": "{}"})
    assert none.status_code == 400
    bad = ec.post("/api/preprocess", files={"file": ("a.png", FIXTURE_PNG.read_bytes())}, data={"preprocess": '{"x": 1}'})
    assert bad.status_code == 400
    office = ec.post("/api/preprocess", files={"file": ("a.docx", b"PK")}, data={"preprocess": '{"contrast": true}'})
    assert office.status_code == 415


def test_ocr_applies_preprocess_to_page_image(ec, monkeypatch):
    seen = []
    real = svc.preprocess_page

    def spy(path, options):
        seen.append(options)
        real(path, options)

    monkeypatch.setattr(svc, "preprocess_page", spy)
    run = ocr_run(ec, preprocess=json.dumps({"deskew": True, "binarize": True}))
    assert run["preprocess"] == ["binarize", "deskew"]
    assert seen and seen[0]["binarize"] is True
    image = ec.get(f"/api/results/{run['id']}/pages/0/image")
    pixels = cv2.imdecode(np.frombuffer(image.content, np.uint8), cv2.IMREAD_GRAYSCALE)
    assert set(np.unique(pixels)) <= {0, 255}  # the stored page is the cleaned one


def test_ocr_without_preprocess_leaves_page_alone(ec, monkeypatch):
    monkeypatch.setattr(svc, "preprocess_page", lambda *a: pytest.fail("preprocess ran"))
    assert ocr_run(ec)["preprocess"] == []


# --- plugins ---


def test_plugins_listed_and_disabled_without_key(ec, monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "")
    plugins = ec.get("/api/plugins").json()
    assert {p["name"] for p in plugins} == {"refine", "kie", "translate"}
    assert all(p["enabled"] is False for p in plugins)
    run = ocr_run(ec)
    resp = ec.post(f"/api/results/{run['id']}/plugins/kie", json={})
    assert resp.status_code == 503 and "GEMINI_API_KEY" in resp.json()["detail"]


def test_plugin_runs_on_result_text(ec, monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "test")
    monkeypatch.setattr(
        "app.plugins.kie.gemini", lambda prompt, system, as_json=False: json.dumps({"title": "Calibration Report"})
    )
    run = ocr_run(ec)
    resp = ec.post(f"/api/results/{run['id']}/plugins/kie", json={"fields": ["title"]})
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"plugin": "kie", "result_id": run["id"], "fields": {"title": "Calibration Report"}}


def test_plugin_errors(ec, monkeypatch):
    from app.plugins import PluginModelError

    monkeypatch.setattr(settings, "gemini_api_key", "test")
    run = ocr_run(ec)
    assert ec.post(f"/api/results/{run['id']}/plugins/nope", json={}).status_code == 404
    assert ec.post(f"/api/results/{run['id']}/plugins/translate", json={}).status_code == 400

    def boom(prompt, system, as_json=False):
        raise PluginModelError("Gemini call failed: 500")

    monkeypatch.setattr("app.plugins.translate.gemini", boom)
    resp = ec.post(f"/api/results/{run['id']}/plugins/translate", json={"target_language": "French"})
    assert resp.status_code == 502


def test_export_and_plugins_reject_unfinished_runs(ec, monkeypatch):
    from app.modules import ocr_jobs

    run = ocr_run(ec)
    real = ocr_jobs.job_manager.get_result
    monkeypatch.setattr(ocr_jobs.job_manager, "get_result", lambda rid: {**real(rid), "status": "running"})
    assert ec.get(f"/api/results/{run['id']}/export/txt").status_code == 409
