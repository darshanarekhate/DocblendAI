"""Tables in typed documents: header kept with every row (Word, PowerPoint, PDF), ruled PDF
tables found and written in page order, merged cells, and table content type."""

from pathlib import Path

from docx import Document as new_docx

from app.models.schemas import ContentType
from app.modules import content_type, text_parser

SALES = [("Month", "Sales", "Customers"), ("Jan", "50,000", "200"), ("Feb", "65,000", "250")]


def _ruled_table_pdf(path: Path, rows, above="Retail store dataset", below="Used to analyze store performance.") -> Path:
    """A one-page PDF: a line of text, a table drawn with ruling lines, and a line of text."""
    col_w, row_h, x0, y_top = 120, 22, 72, 650
    ops = [f"BT /F1 11 Tf 72 700 Td ({above}) Tj ET"]
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            ops.append(f"BT /F1 10 Tf {x0 + c * col_w + 6} {y_top - (r + 1) * row_h + 7} Td ({cell}) Tj ET")
    n_rows, n_cols = len(rows), len(rows[0])
    for r in range(n_rows + 1):  # horizontal rules
        y = y_top - r * row_h
        ops.append(f"{x0} {y} m {x0 + n_cols * col_w} {y} l S")
    for c in range(n_cols + 1):  # vertical rules
        x = x0 + c * col_w
        ops.append(f"{x} {y_top} m {x} {y_top - n_rows * row_h} l S")
    ops.append(f"BT /F1 11 Tf 72 {y_top - n_rows * row_h - 40} Td ({below}) Tj ET")
    stream = "\n".join(ops).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents 5 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)
    return path


# --- PDF -----------------------------------------------------------------------------------


def test_pdf_table_rows_keep_their_header(tmp_path) -> None:
    [(text, conf)] = text_parser.parse(str(_ruled_table_pdf(tmp_path / "sales.pdf", SALES)))
    lines = text.splitlines()
    assert conf == 1.0
    assert "Table: Month | Sales | Customers" in lines
    assert "Month: Feb | Sales: 65,000 | Customers: 250" in lines


def test_pdf_table_stays_between_the_text_above_and_below(tmp_path) -> None:
    [(text, _)] = text_parser.parse(str(_ruled_table_pdf(tmp_path / "sales.pdf", SALES)))
    lines = text.splitlines()
    assert lines[0] == "Retail store dataset"
    assert lines[-1] == "Used to analyze store performance."
    assert sum("Month:" in line for line in lines) == 2  # each row once, not also as loose text


def test_pdf_without_tables_is_unchanged(pdf_file) -> None:
    [(text, _)] = text_parser.parse(str(pdf_file(["Plain notes.\nNo table here."])))
    assert text == "Plain notes.\nNo table here."


# --- Word ----------------------------------------------------------------------------------


def _docx_table(path: Path, rows) -> Path:
    doc = new_docx()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(path)
    return path


def test_equal_values_in_neighbouring_cells_are_both_kept(tmp_path) -> None:
    rows = [("Feature", "Bar chart", "Pie chart"), ("Shows parts", "Yes", "Yes")]
    [(text, _)] = text_parser.parse(str(_docx_table(tmp_path / "t.docx", rows)))
    assert "Feature: Shows parts | Bar chart: Yes | Pie chart: Yes" in text.splitlines()


def test_merged_cell_is_written_once(tmp_path) -> None:
    doc = new_docx()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Week 1 summary"
    for cell, value in zip(table.rows[1].cells, ("Mon", "Tue", "Wed")):
        cell.text = value
    doc.save(tmp_path / "merged.docx")
    [(text, _)] = text_parser.parse(str(tmp_path / "merged.docx"))
    assert text.splitlines() == ["Week 1 summary", "Mon | Tue | Wed"]  # no usable header: plain rows


def test_cell_text_on_several_lines_becomes_one_line(tmp_path) -> None:
    rows = [("Term", "Meaning"), ("2NF", "No partial\ndependencies")]
    [(text, _)] = text_parser.parse(str(_docx_table(tmp_path / "t.docx", rows)))
    assert "Term: 2NF | Meaning: No partial dependencies" in text.splitlines()


# --- Module 4 ------------------------------------------------------------------------------


def test_two_column_table_is_labelled_table() -> None:
    text = "Table: Name | Marks\nName: Asha | Marks: 91\nName: Ravi | Marks: 84"
    assert content_type.classify(text) is ContentType.TABLE


def test_prose_with_a_colon_is_still_a_paragraph() -> None:
    text = "Note: normal forms reduce redundancy.\nThey are applied step by step.\nEach removes one kind of dependency."
    assert content_type.classify(text) is ContentType.PARAGRAPH


# --- chunking: tables get chunks of their own ------------------------------------------------

from app.modules import chunker  # noqa: E402

PAGE = (
    "Unit 5 covers visualization of a store's data and how to choose charts for it.\n"
    "Dataset contains:\n"
    "Table: Month | Sales | Customers\n"
    "Month: Jan | Sales: 50,000 | Customers: 200\n"
    "Month: Feb | Sales: 65,000 | Customers: 250\n"
    "This dataset is used to analyze store performance."
)


def test_a_table_is_its_own_chunk_with_its_caption() -> None:
    chunks = chunker.chunk_pages_numbered("d", [(PAGE, 1.0)])
    texts = [c.text for _, c in chunks]
    table = next(t for t in texts if "Customers: 250" in t)
    assert table.splitlines() == PAGE.splitlines()[1:5]  # caption + header + rows, no prose
    assert any(t.startswith("Unit 5 covers") for t in texts) and any("analyze store" in t for t in texts)
    loc = next(loc for loc, c in chunks if c.text == table)
    assert (loc.page, loc.first_line, loc.last_line) == (1, 2, 5)


def test_a_long_table_repeats_its_header_in_every_part() -> None:
    rows = "\n".join(f"Roll: {n} | Name: Student {n} | Marks: {50 + n}" for n in range(1, 41))
    chunks = chunker.chunk_pages("d", [("Table: Roll | Name | Marks\n" + rows, 1.0)], chunk_size=300, overlap=50)
    assert len(chunks) > 3
    assert all(c.text.startswith("Table: Roll | Name | Marks") for c in chunks)
    all_rows = [line for c in chunks for line in c.text.splitlines() if not line.startswith("Table:")]
    assert all_rows == rows.splitlines()  # every row exactly once, in order


def test_pages_without_tables_are_chunked_as_before() -> None:
    text = " ".join(f"word{n}" for n in range(400))
    assert [c.text for c in chunker.chunk_pages("d", [(text, 1.0)])] == [p for p, _, _ in chunker._split(text, 1000, 200)]
