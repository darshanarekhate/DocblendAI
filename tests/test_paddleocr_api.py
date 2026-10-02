"""Experience Center HTTP API (/api/*) with the Paddle engine faked: contract §3-§4."""

import io
import json
import time
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from app.calibration import calibrate as calibrate_module
from app.config import settings
from app.db.models import OCRRunORM
from app.modules import ocr_jobs
from app.modules import paddleocr_service as svc
from app.routers import paddleocr as router_module
from tests.conftest import make_pdf
from tests.paddle_fakes import FakeFactory

FIXTURE_PNG = Path(__file__).parent / "fixtures" / "ppstructurev3_page.png"


@pytest.fixture
def ec(client, db_session_factory, monkeypatch):
    """Client + fake Paddle models; the job manager writes to the client's temp SQLite."""
    factory = FakeFactory()
    monkeypatch.setattr(svc, "_construct_model", factory)
    monkeypatch.setattr(svc.PaddleOCRService, "installed", staticmethod(lambda: True))
    monkeypatch.setattr(settings, "extraction_engine", "paddle")
    ocr_jobs.job_manager.configure(db_session_factory)
    client.factory = factory
    return client


def png_bytes(size=(120, 80)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, "white").save(buf, format="PNG")
    return buf.getvalue()


def post_ocr(client, name="page.png", data=None, **form):
    form.setdefault("wait", "true")
    return client.post("/api/ocr", files={"file": (name, data if data is not None else png_bytes())}, data=form)


def poll(client, run_id, timeout=10.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        body = client.get(f"/api/results/{run_id}").json()
        if body["status"] in ("done", "error"):
            return body
        time.sleep(0.02)
    raise AssertionError("job did not finish")


def test_health_shape(ec):
    body = ec.get("/api/health").json()
    assert body["status"] == "ok"
    assert set(body["engine"]) == {"paddleocr", "paddlepaddle", "ocr_version", "device"}
    assert set(body["pipelines"]) == {"ocr", "structure", "vl", "office"}
    assert body["pipelines"]["ocr"]["available"] is True
    assert body["pipelines"]["vl"] == {"available": False, "message": svc.VL_DISABLED_MESSAGE}
    assert body["calibration"] == {"status": "uncalibrated", "method": None}
    assert body["review_threshold"] == settings.review_threshold
    assert body["max_upload_mb"] == settings.paddle_max_upload_mb
    assert {"code": "en", "name": "English"} in body["languages"]
    assert ec.factory.built == []  # health never loads a model
    assert ec.get("/health").json() == {"status": "ok"}  # the original route is untouched


def test_ocr_png_wait_returns_full_result(ec, db_session_factory):
    resp = post_ocr(ec, review_threshold="0.98")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["status"] == "done" and body["progress"] == 1.0 and body["error"] is None
    assert body["pipeline"] == "ocr" and body["language"] == "en" and body["filename"] == "page.png"
    assert body["calibration"] == {"status": "uncalibrated", "method": None, "review_threshold": 0.98}
    assert body["summary"]["page_count"] == 1 and body["summary"]["line_count"] == 15
    assert body["summary"]["needs_review"] == 1
    assert body["markdown"].startswith("Calibration Report\n")
    (page,) = body["pages"]
    assert (page["width"], page["height"]) == (120, 80)
    assert page["image_url"] == f"/api/results/{body['id']}/pages/0/image"
    assert page["lines"][0]["words"][0]["text"] == "Calibration"
    image = ec.get(page["image_url"])
    assert image.status_code == 200 and image.headers["content-type"] == "image/png"
    assert ec.get(f"/api/results/{body['id']}").json() == body
    with db_session_factory() as db:
        row = db.get(OCRRunORM, body["id"])
        assert row.status == "done" and row.page_count == 1 and "Temperature" in row.full_text


def test_ocr_two_page_pdf(ec, monkeypatch):
    monkeypatch.setattr(settings, "paddle_pdf_dpi", 72)
    resp = post_ocr(ec, "scan.pdf", make_pdf(["first page", "second page"]))
    body = resp.json()
    assert resp.status_code == 200 and body["summary"]["page_count"] == 2
    assert [p["index"] for p in body["pages"]] == [0, 1]
    assert body["pages"][1]["width"] == 612
    assert ec.get(f"/api/results/{body['id']}/pages/1/image").status_code == 200
    assert ec.get(f"/api/results/{body['id']}/pages/2/image").status_code == 404


def test_async_job_and_poll(ec):
    resp = post_ocr(ec, wait="false")
    assert resp.status_code == 202
    accepted = resp.json()
    assert set(accepted) == {"id", "status"} and accepted["status"] in ("queued", "running", "done")
    body = poll(ec, accepted["id"])
    assert body["status"] == "done" and body["pages"]


def test_upload_temp_files_are_removed(ec, tmp_path, monkeypatch):
    import tempfile

    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path / "tmp"))
    (tmp_path / "tmp").mkdir()
    post_ocr(ec)
    post_ocr(ec, data=b"not a png")
    assert list((tmp_path / "tmp").iterdir()) == []


def test_parse_structure_and_language(ec):
    resp = ec.post(
        "/api/parse", files={"file": ("p.png", FIXTURE_PNG.read_bytes())},
        data={"wait": "true", "language": "fr", "preprocess": '{"deskew": true}'},
    )
    body = resp.json()
    assert resp.status_code == 200 and body["pipeline"] == "structure" and body["language"] == "fr"
    page = body["pages"][0]
    assert [b["type"] for b in page["blocks"]] == ["paragraph_title", "text", "table"]
    assert page["tables"][0]["rows"][1] == ["Raw", "0.081", "0.142"]
    assert body["markdown"].startswith("## Calibration Report")
    assert ec.factory.built == [("structure", "fr")]


def test_parse_vl_disabled_is_503(ec):
    resp = ec.post("/api/parse", files={"file": ("p.png", png_bytes())}, data={"pipeline": "vl"})
    assert resp.status_code == 503 and "PADDLE_ENABLE_VL" in resp.json()["detail"]


def test_model_load_failure_is_503_when_waiting(ec):
    ec.factory.fail = RuntimeError("weights missing")
    resp = post_ocr(ec)
    assert resp.status_code == 503 and "weights missing" in resp.json()["detail"]
    (item,) = ec.get("/api/results").json()
    assert item["status"] == "error"


def test_office_file_runs_office_pipeline(ec, monkeypatch):
    from docx import Document as WordDocument

    monkeypatch.setattr(svc.PaddleOCRService, "convert_office", lambda self, p: "# Notes\n\n| a | b |\n|---|---|\n| 1 | 2 |")
    doc = WordDocument()
    doc.add_heading("Notes", 1)
    doc.add_paragraph("Normal forms remove redundancy.")
    buf = io.BytesIO()
    doc.save(buf)
    body = post_ocr(ec, "notes.docx", buf.getvalue()).json()
    assert body["pipeline"] == "office" and body["status"] == "done"
    assert body["pages"][0]["image_url"] is None and body["pages"][0]["tables"][0]["rows"][1] == ["1", "2"]
    assert body["markdown"].startswith("# Notes")
    # Exact text (Module 2's typed path) gives the lines; typed text needs no calibration.
    assert [l["text"] for l in body["pages"][0]["lines"]] == ["Notes", "Normal forms remove redundancy."]
    assert body["format_type"] == "typed" and body["engine"] == "text"
    assert body["calibration"]["status"] == "calibrated"


def test_history_search_limit_and_order(ec):
    first = post_ocr(ec, "alpha.png").json()
    second = post_ocr(ec, "beta.png").json()
    items = ec.get("/api/results").json()
    assert [i["id"] for i in items][:2] in ([second["id"], first["id"]], [first["id"], second["id"]])
    assert set(items[0]) == {
        "id", "filename", "pipeline", "language", "status", "created_at", "page_count",
        "mean_confidence", "processing_time_ms",
    }
    assert [i["filename"] for i in ec.get("/api/results", params={"q": "alpha"}).json()] == ["alpha.png"]
    assert len(ec.get("/api/results", params={"q": "isotonic"}).json()) == 2  # matches recognised text
    assert ec.get("/api/results", params={"q": "100%_nothing"}).json() == []
    assert len(ec.get("/api/results", params={"limit": 1}).json()) == 1
    assert ec.get("/api/results", params={"limit": 0}).status_code == 422


def test_patch_lines_marks_edits(ec):
    run = post_ocr(ec, review_threshold="0.98").json()
    resp = ec.patch(f"/api/results/{run['id']}/lines", json=[{"page": 0, "line": 0, "text": "Calibration Report!"}])
    assert resp.status_code == 200, resp.text
    body = resp.json()
    line = body["pages"][0]["lines"][0]
    assert line["text"] == "Calibration Report!" and line["edited"] is True and line["needs_review"] is False
    assert body["pages"][0]["lines"][1]["edited"] is False
    assert body["summary"]["needs_review"] == 0
    assert body["markdown"].startswith("Calibration Report!\n")
    persisted = ec.get(f"/api/results/{run['id']}").json()
    assert persisted.pop("active_job") is None
    assert persisted == body  # persisted
    assert len(ec.get("/api/results", params={"q": "Report!"}).json()) == 1

    bad = ec.patch(f"/api/results/{run['id']}/lines", json=[{"page": 0, "line": 99, "text": "x"}])
    assert bad.status_code == 400
    assert ec.patch(f"/api/results/{run['id']}/lines", json=[{"page": "0", "line": 0}]).status_code == 400
    assert ec.patch("/api/results/nope/lines", json=[]).status_code == 404


def test_delete_removes_row_and_images(ec):
    run = post_ocr(ec).json()
    run_dir = settings.paddle_runs_dir / run["id"]
    assert (run_dir / "page-0.png").is_file()
    assert ec.delete(f"/api/results/{run['id']}").json() == {"id": run["id"], "deleted": True}
    assert not run_dir.exists()
    assert ec.get(f"/api/results/{run['id']}").status_code == 404
    assert ec.delete(f"/api/results/{run['id']}").status_code == 404


def test_interrupted_runs_report_error(ec, db_session_factory):
    with db_session_factory() as db:
        db.add(OCRRunORM(id="a" * 32, filename="x.png", pipeline="ocr", language="en", status="running",
                         created_at="2026-01-01T00:00:00+00:00", full_text="",
                         result_json=json.dumps({"id": "a" * 32, "status": "running", "pages": []})))
        db.commit()
    body = ec.get(f"/api/results/{'a' * 32}").json()
    assert body["status"] == "error" and "Interrupted" in body["error"]
    assert ec.get("/api/results").json()[0]["status"] == "error"


def test_unknown_result_404(ec):
    assert ec.get("/api/results/doesnotexist").status_code == 404
    assert ec.get("/api/results/doesnotexist/pages/0/image").status_code == 404


# -- validation --------------------------------------------------------------------------------


def test_too_large_is_413(ec, monkeypatch):
    monkeypatch.setattr(settings, "paddle_max_upload_mb", 0)
    resp = post_ocr(ec)
    assert resp.status_code == 413 and "MB" in resp.json()["detail"]


@pytest.mark.parametrize(("name", "data"), [("notes.rtf", b"hello"), ("fake.png", b"hello"), ("fake.pdf", b"%PDF-1.4 junk")])
def test_bad_type_is_415(ec, name, data):
    assert post_ocr(ec, name, data).status_code == 415
    assert ec.get("/api/results").json() == []


@pytest.mark.parametrize(
    "form",
    [
        {"language": "xx"},
        {"review_threshold": "1.5"},
        {"review_threshold": "abc"},
        {"preprocess": "[1]"},
        {"preprocess": "{not json"},
        {"preprocess": '{"sharpen": true}'},
        {"preprocess": '{"deskew": "yes"}'},
    ],
)
def test_bad_form_values_are_400(ec, form):
    resp = post_ocr(ec, **form)
    assert resp.status_code == 400 and resp.json()["detail"]


def test_bad_parse_pipeline_is_400(ec):
    resp = ec.post("/api/parse", files={"file": ("p.png", png_bytes())}, data={"pipeline": "magic"})
    assert resp.status_code == 400


# -- calibration -------------------------------------------------------------------------------


def dataset_zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


@pytest.fixture
def fake_calibration(monkeypatch):
    calls = []

    def run_calibration(dataset, ocr_fn, out_dir=None, match="cer", cer_threshold=None, progress=None):
        dataset = Path(dataset)
        calls.append({"dataset": dataset, "files": sorted(p.name for p in dataset.iterdir()),
                      "match": match, "cer_threshold": cer_threshold, "ocr_fn": ocr_fn})
        if progress:
            progress(1, 2)
            progress(2, 2)
        out = Path(out_dir or settings.paddle_calibration_dir)
        out.mkdir(parents=True, exist_ok=True)
        report = {"method": "isotonic", "n_lines": 2, "diagram": "reliability_diagram.png"}
        (out / "report.json").write_text(json.dumps(report))
        (out / "reliability_diagram.png").write_bytes(png_bytes((10, 10)))
        return report

    monkeypatch.setattr(calibrate_module, "run_calibration", run_calibration)
    return calls


def test_calibrate_with_zip(ec, fake_calibration):
    assert ec.get("/api/calibration").json() == {"status": "uncalibrated"}
    assert ec.get("/api/calibration/diagram.png").status_code == 404
    data = dataset_zip({"set/labels.csv": b"image,text\na.png,hello\n", "set/a.png": png_bytes()})
    resp = ec.post("/api/calibrate", files={"file": ("set.zip", data)}, data={"match": "exact", "cer_threshold": "0.2"})
    assert resp.status_code == 202
    body = poll(ec, resp.json()["id"])
    assert body["status"] == "done" and body["pipeline"] == "calibration"
    assert body["report"]["method"] == "isotonic"
    (call,) = fake_calibration
    assert call["files"] == ["a.png", "labels.csv"] and call["match"] == "exact" and call["cer_threshold"] == 0.2
    assert call["ocr_fn"] == svc.paddleocr_service.recognize_lines
    assert not call["dataset"].exists()  # extracted temp dir cleaned up
    report = ec.get("/api/calibration").json()
    assert report["method"] == "isotonic"
    assert ec.get("/api/calibration/diagram.png").headers["content-type"] == "image/png"


def test_calibrate_with_sample_and_wait(ec, fake_calibration, tmp_path, monkeypatch):
    sample = tmp_path / "calibration_sample"
    monkeypatch.setattr(router_module, "SAMPLE_DATASET", sample)
    assert ec.post("/api/calibrate", data={"use_sample": "true"}).status_code == 503
    sample.mkdir()
    (sample / "labels.csv").write_text("image,text\n")
    resp = ec.post("/api/calibrate", data={"use_sample": "true", "wait": "true"})
    assert resp.status_code == 200 and resp.json()["status"] == "done"
    assert fake_calibration[0]["dataset"] == sample and fake_calibration[0]["match"] == "cer"
    assert fake_calibration[0]["cer_threshold"] == settings.calibration_cer_threshold


def test_calibration_failure_is_reported(ec, monkeypatch, tmp_path):
    monkeypatch.setattr(router_module, "SAMPLE_DATASET", tmp_path)
    (tmp_path / "labels.csv").write_text("image,text\n")

    def boom(*args, **kwargs):
        raise NotImplementedError("calibration is not implemented yet")

    monkeypatch.setattr(calibrate_module, "run_calibration", boom)
    body = ec.post("/api/calibrate", data={"use_sample": "true", "wait": "true"}).json()
    assert body["status"] == "error" and "not implemented" in body["error"]


@pytest.mark.parametrize(
    ("entries", "code"),
    [
        ({"../evil.txt": b"x", "labels.csv": b"image,text\n"}, 400),
        ({"/abs.txt": b"x", "labels.csv": b"image,text\n"}, 400),
        ({"a.png": b"x"}, 400),  # no labels.csv
        ({"labels.csv": b"file,label\n"}, 400),  # wrong columns
    ],
)
def test_calibrate_rejects_bad_zips(ec, fake_calibration, entries, code):
    resp = ec.post("/api/calibrate", files={"file": ("d.zip", dataset_zip(entries))})
    assert resp.status_code == code, resp.text
    assert fake_calibration == []


def test_calibrate_form_validation(ec, fake_calibration, monkeypatch):
    assert ec.post("/api/calibrate").status_code == 400
    assert ec.post("/api/calibrate", data={"use_sample": "true", "match": "fuzzy"}).status_code == 400
    assert ec.post("/api/calibrate", data={"use_sample": "true", "cer_threshold": "2"}).status_code == 400
    assert ec.post("/api/calibrate", files={"file": ("d.tar", b"x")}).status_code == 415
    assert ec.post("/api/calibrate", files={"file": ("d.zip", b"not a zip")}).status_code == 415
    monkeypatch.setattr(router_module, "_zip_max_uncompressed", lambda: 10)
    big = dataset_zip({"labels.csv": b"image,text\n" + b"x" * 100})
    assert ec.post("/api/calibrate", files={"file": ("d.zip", big)}).status_code == 413
