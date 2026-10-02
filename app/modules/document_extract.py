"""Module 2 — Format Detection & Text Extraction, shared by the QA page and the Experience Center.

Responsibility: read one uploaded file into Experience Center pages
(docs/experience_center_contract.md §4: lines with boxes, words, raw and calibrated
confidence, layout blocks, tables, Markdown), picking the extraction method per format:

- typed (PDF with a text layer, .docx, .pptx, .xlsx, .txt): exact text, no recognition.
  PDFs also get their line, word and table boxes from pdfplumber, and a page image.
- scanned (images, PDFs without text): PaddleOCR text lines or PP-StructureV3 layout
  (blocks, tables, formulas), confidence calibrated by app/calibration (Module 3).
- handwritten: PaddleOCR finds the lines, TrOCR reads each line box (htr_extractor),
  confidence calibrated by confidence_capture's handwritten table.

Typed vs not is decided by the text layer (format_detection.has_text_layer). Scanned vs
handwritten, unless the user said which: Paddle's confidence; when it is low, TrOCR reads a
sample of the same line boxes and the engine with the higher calibrated confidence wins
(the same rule as format_detection.detect_format, without a second OCR pass).

With settings.extraction_engine = "classic" (or when PaddleOCR is not installed) scanned and
handwritten pages go through format_detection's Tesseract/docTR and TrOCR path, one entry
per page; their lines are the page's text lines, without boxes.

Every page also keeps "text": the exact text extracted for typed and classic pages, which the
QA index (qa_index.py) chunks as-is until a line is edited.

Uses from schemas.py: Document, FormatType.
"""

from __future__ import annotations

import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.config import settings
from app.models.schemas import Document, FormatType
from app.modules import confidence_capture
from app.modules import paddleocr_service as svc
from app.modules.file_types import FileKind, kind_of

logger = logging.getLogger(__name__)

ProgressFn = Callable[[float, str], None]
ImageUrlFn = Callable[[int], str]

# Office kinds (incl. .xlsx, which only the Experience Center converted before) are typed.
TYPED_SUFFIXES = frozenset({".docx", ".pptx", ".xlsx", ".txt"})
# Pages sampled with TrOCR to decide scanned vs handwritten (same as format_detection).
HTR_SAMPLE_PAGES = 2
POINTS_PER_INCH = 72


class ExtractionError(ValueError):
    """The file cannot be read as the kind its extension claims (-> HTTP 415/422)."""


@dataclass
class Extraction:
    format_type: FormatType
    pages: list[dict[str, Any]]
    page_count: int  # pages in the source (pages may hold fewer: PADDLE_MAX_PAGES)
    engine: str  # what read the text: text | paddleocr | pp-structurev3 | paddleocr+trocr | tesseract | trocr
    notes: list[str] = field(default_factory=list)


def engine_mode() -> str:
    """"paddle" or "classic" for this server (settings.extraction_engine, "auto" = paddle if installed)."""
    mode = (settings.extraction_engine or "auto").lower()
    if mode == "auto":
        return "paddle" if svc.PaddleOCRService.installed() else "classic"
    return "paddle" if mode == "paddle" else "classic"


def _line(page_index: int, i: int, text: str, raw: float, calibrated: float, source: str, **extra: Any) -> dict[str, Any]:
    line = {
        "id": f"p{page_index}-l{i}", "text": text, "polygon": [], "box": None,
        "raw_confidence": round(raw, 6), "calibrated_confidence": round(calibrated, 6),
        "needs_review": False, "edited": False, "source": source, "words": [],
    }
    line.update(extra)
    return line


def _text_lines(page_index: int, text: str, raw: float, calibrated: float, source: str) -> list[dict[str, Any]]:
    """A page's text as lines without boxes (typed Office/txt files, classic OCR/HTR)."""
    rows = [row.rstrip() for row in text.splitlines() if row.strip()]
    return [_line(page_index, i, row, raw, calibrated, source) for i, row in enumerate(rows)]


def _page(index: int, text: str | None = None, image_url: str | None = None, width=None, height=None) -> dict[str, Any]:
    return {
        "index": index, "width": width, "height": height, "image_url": image_url,
        "lines": [], "markdown": "", "blocks": [], "tables": [], "text": text,
    }


def _report(progress: ProgressFn | None, fraction: float, message: str) -> None:
    if progress:
        progress(fraction, message)


# ---------------------------------------------------------------------------------------------
# typed files: exact text
# ---------------------------------------------------------------------------------------------


def _office_markdown(path: Path, pages_text: list[str]) -> str:
    """Structured Markdown for an Office/txt file: doc2md (Paddle mode), else the text itself."""
    if path.suffix.lower() == ".txt":
        return "\n\n".join(pages_text)
    if engine_mode() == "paddle":
        try:
            return svc.paddleocr_service.convert_office(path)
        except svc.PaddleOCRUnavailableError as exc:
            logger.warning("Office conversion unavailable (%s); using the plain text", exc)
        except Exception as exc:  # a converter bug must not lose the exact text
            logger.warning("Office conversion of %s failed (%s); using the plain text", path.name, exc)
    return "\n\n".join(pages_text)


def _parse_xlsx(path: Path) -> list[tuple[str, float]]:
    """One entry per sheet, cells separated like text_parser's table rows (two spaces)."""
    from openpyxl import load_workbook

    from app.modules.text_parser import CELL_SEPARATOR, TYPED_CONFIDENCE

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"not a readable Excel workbook: {exc}") from exc
    pages = []
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(CELL_SEPARATOR.join(c for c in cells if c))
            pages.append(("\n".join(rows), TYPED_CONFIDENCE))
    finally:
        book.close()
    return pages


def _typed_office(path: Path) -> Extraction:
    from app.modules import text_parser
    from app.modules.file_types import UnreadableFileError

    try:
        parsed = _parse_xlsx(path) if path.suffix.lower() == ".xlsx" else text_parser.parse(str(path))
    except UnreadableFileError as exc:
        raise ExtractionError(str(exc)) from exc
    texts = [text for text, _ in parsed] or [""]
    markdown = _office_markdown(path, texts)
    pages = []
    for i, text in enumerate(texts):
        page = _page(i, text)
        page["lines"] = _text_lines(i, text, 1.0, 1.0, "text")
        page["markdown"] = text
        pages.append(page)
    # One converted document: its Markdown (and tables) go on the first page when it is the only one.
    if len(pages) == 1:
        pages[0]["markdown"] = markdown.strip()
        pages[0]["tables"] = svc.tables_from_markdown(markdown)
    if path.suffix.lower() == ".txt":
        for page in pages:
            page["preformatted"] = True  # keep the file's own spacing in the layout view
    return Extraction(FormatType.TYPED, pages, len(pages), "text")


def _pdf_table(table: Any, scale: float) -> dict[str, Any]:
    """pdfplumber table -> contract table (None cells continue a merged cell to their left)."""
    from html import escape

    rows = table.extract() or []
    html_rows = []
    for row in rows:
        cells: list[list[Any]] = []
        for value in row:
            if value is None and cells:
                cells[-1][1] += 1
            else:
                cells.append([value or "", 1])
        html_rows.append("".join(
            f"<td{f' colspan={span}' if span > 1 else ''}>{escape(' '.join(str(v).split()))}</td>" for v, span in cells
        ))
    clean_rows = [[" ".join(str(c or "").split()) for c in row] for row in rows]
    return {
        "box": [round(v * scale, 2) for v in table.bbox],
        "html": "<table>" + "".join(f"<tr>{r}</tr>" for r in html_rows) + "</table>",
        "rows": clean_rows,
        "cells": [[round(v * scale, 2) for v in cell] for cell in (table.cells or [])],
    }


def _typed_pdf(path: Path, pages_dir: Path, image_url: ImageUrlFn | None, progress: ProgressFn | None) -> Extraction:
    import pdfplumber

    image_paths, total = svc.render_pages(path, "pdf", pages_dir)
    scale = settings.paddle_pdf_dpi / POINTS_PER_INCH
    pages = []
    with pdfplumber.open(path) as pdf:
        n = min(len(pdf.pages), len(image_paths))
        for i in range(n):
            _report(progress, i / max(1, n), f"Page {i + 1} of {n}")
            pp = pdf.pages[i]
            text = (pp.extract_text() or "").strip()  # exactly what text_parser gives the QA index
            width, height = svc._image_size(image_paths[i])
            page = _page(i, text, image_url(i) if image_url else None, width, height)
            words = pp.extract_words() or []
            lines = []
            for j, row in enumerate(pp.extract_text_lines(return_chars=False) or []):
                box = [round(row["x0"] * scale, 2), round(row["top"] * scale, 2),
                       round(row["x1"] * scale, 2), round(row["bottom"] * scale, 2)]
                mid = lambda w: (w["top"] + w["bottom"]) / 2  # noqa: E731
                line_words = [
                    {"text": w["text"], "box": [round(w["x0"] * scale, 2), round(w["top"] * scale, 2),
                                                round(w["x1"] * scale, 2), round(w["bottom"] * scale, 2)]}
                    for w in words
                    if row["top"] - 0.5 <= mid(w) <= row["bottom"] + 0.5 and row["x0"] - 0.5 <= w["x0"] <= row["x1"] + 0.5
                ]
                polygon = [[box[0], box[1]], [box[2], box[1]], [box[2], box[3]], [box[0], box[3]]]
                lines.append(_line(i, j, row["text"], 1.0, 1.0, "text", box=box, polygon=polygon, words=line_words))
            page["lines"] = lines
            page["tables"] = [_pdf_table(t, scale) for t in pp.find_tables()]
            page["blocks"] = [{"type": "table", "box": t["box"], "content": t["html"]} for t in page["tables"]]
            page["markdown"] = text
            pages.append(page)
    notes = []
    if total > len(image_paths):
        notes.append(f"Only the first {len(image_paths)} of {total} pages were processed (PADDLE_MAX_PAGES).")
    return Extraction(FormatType.TYPED, pages, total, "text", notes)


# ---------------------------------------------------------------------------------------------
# scanned / handwritten: PaddleOCR (+ TrOCR in Paddle's line boxes)
# ---------------------------------------------------------------------------------------------


def _weighted_conf(lines: list[dict[str, Any]], key: str = "calibrated_confidence") -> float:
    chars = sum(len(line["text"]) for line in lines)
    return sum(line[key] * len(line["text"]) for line in lines) / chars if chars else 0.0


def _calibrate_paddle(pages: list[dict[str, Any]]) -> None:
    from app.calibration import calibrate_confidences

    lines = svc.all_lines(pages)
    if not lines:
        return
    calibrated, _ = calibrate_confidences([line["raw_confidence"] for line in lines])
    for line, value in zip(lines, calibrated):
        line["calibrated_confidence"] = round(svc._clamp01(value), 6)
        line["source"] = "paddleocr"


def _htr_page(page: dict[str, Any], image_path: Path) -> None:
    """Replace each line's text with TrOCR's reading of the line's box (handwriting)."""
    from PIL import Image

    from app.modules import htr_extractor, line_segmentation
    from app.modules.pdf_render import denoise

    lines = [line for line in page["lines"] if line.get("box")]
    if not lines:
        return
    with Image.open(image_path) as img:
        gray = denoise(img.convert("L"))
    if settings.remove_ruled_lines:
        gray = line_segmentation.remove_ruled_lines(gray)
    crops = []
    for line in lines:
        x0, y0, x1, y1 = line["box"]
        pad = 0.15 * (y1 - y0)
        crops.append(gray.crop((max(0, int(x0 - pad)), max(0, int(y0 - pad)),
                                min(gray.width, int(x1 + pad)), min(gray.height, int(y1 + pad)))))
    for line, (text, conf) in zip(lines, htr_extractor.recognize_lines(crops)):
        line["paddle_text"] = line["text"]
        line["text"] = text
        line["raw_confidence"] = round(svc._clamp01(conf), 6)
        line["calibrated_confidence"] = round(confidence_capture.calibrate(conf, FormatType.HANDWRITTEN), 6)
        line["source"] = "trocr"
        line["words"] = []  # Paddle's word boxes belong to its own reading, not TrOCR's
    page["lines"] = [line for line in page["lines"] if line["text"].strip()]
    if not page.get("blocks"):
        page["markdown"] = svc.lines_markdown(page["lines"])


def _paddle_pages(
    path: Path, kind: str, pipeline: str, lang: str, pages_dir: Path, image_url: ImageUrlFn | None,
    progress: ProgressFn | None, preprocess: dict[str, bool] | None,
) -> tuple[list[dict[str, Any]], list[Path], int]:
    from app.modules import preprocess as preprocess_steps

    pipeline = pipeline if pipeline in ("ocr", "structure", "vl") else "ocr"
    svc.paddleocr_service.get_model(pipeline, lang, progress)  # load first: its time is not "page 1"
    _report(progress, 0.0, "Preparing pages")
    image_paths, total = svc.render_pages(path, kind, pages_dir)
    if preprocess_steps.any_enabled(preprocess):
        _report(progress, 0.0, "Cleaning up page images")
        for image_path in image_paths:
            svc.preprocess_page(image_path, preprocess)
    pages = []
    n = len(image_paths)
    for i, image_path in enumerate(image_paths):
        _report(progress, i / max(1, n), f"Page {i + 1} of {n}")
        url = image_url(i) if image_url else None
        if pipeline == "ocr":
            page = svc.paddleocr_service.ocr_page(image_path, i, lang, url)
        else:
            page = svc.paddleocr_service.structure_page(image_path, i, lang, pipeline, url)
        page.setdefault("text", None)
        pages.append(page)
    _calibrate_paddle(pages)
    return pages, image_paths, total


def _decide_handwritten(pages: list[dict[str, Any]], image_paths: list[Path], notes: list[str]) -> bool:
    """Low Paddle confidence: does TrOCR read a sample of the same line boxes better?"""
    from app.modules import htr_extractor

    if _weighted_conf(svc.all_lines(pages)) >= settings.scanned_min_ocr_conf:
        return False
    sample = [dict(p, lines=[dict(l) for l in p["lines"]]) for p in pages[:HTR_SAMPLE_PAGES]]
    paddle_conf = _weighted_conf(svc.all_lines(sample))
    try:
        for page, image_path in zip(sample, image_paths):
            _htr_page(page, image_path)
    except htr_extractor.HTRUnavailableError as exc:
        notes.append(f"Handwriting recognition is unavailable ({exc}); read as print.")
        return False
    return _weighted_conf(svc.all_lines(sample)) > paddle_conf


def _paddle_extract(
    path: Path, kind: str, hint: FormatType | None, pipeline: str, lang: str, pages_dir: Path,
    image_url: ImageUrlFn | None, progress: ProgressFn | None, preprocess: dict[str, bool] | None,
) -> Extraction:
    pages, image_paths, total = _paddle_pages(path, kind, pipeline, lang, pages_dir, image_url, progress, preprocess)
    notes = []
    if total > len(image_paths):
        notes.append(f"Only the first {len(image_paths)} of {total} pages were processed (PADDLE_MAX_PAGES).")
    handwritten = hint is FormatType.HANDWRITTEN if hint is not None else _decide_handwritten(pages, image_paths, notes)
    engine = "pp-structurev3" if pipeline == "structure" else "paddleocr-vl" if pipeline == "vl" else "paddleocr"
    if not handwritten:
        return Extraction(FormatType.SCANNED, pages, total, engine, notes)
    n = len(pages)
    for i, (page, image_path) in enumerate(zip(pages, image_paths)):
        _report(progress, i / max(1, n), f"Reading handwriting, page {i + 1} of {n}")
        _htr_page(page, image_path)
    return Extraction(FormatType.HANDWRITTEN, pages, total, f"{engine}+trocr", notes)


# ---------------------------------------------------------------------------------------------
# classic engine (Tesseract / docTR / TrOCR on whole pages)
# ---------------------------------------------------------------------------------------------


def _classic_extract(
    path: Path, kind: str, hint: FormatType | None, pages_dir: Path, image_url: ImageUrlFn | None,
    progress: ProgressFn | None,
) -> Extraction:
    from app.modules import format_detection

    _report(progress, 0.0, "Detecting the document format")
    page_count = format_detection.page_count(str(path))
    format_type = format_detection.resolve_format(str(path), hint)
    document = Document(doc_id="extract", file_path=str(path), format_type=format_type, page_count=page_count)
    _report(progress, 0.1, "Reading the pages")
    extracted = format_detection.extract(document)
    image_paths: list[Path] = []
    try:
        image_paths, _ = svc.render_pages(path, kind, pages_dir)
    except svc.UnreadableInputError as exc:  # display only: the text was already read
        logger.warning("could not render page images of %s: %s", path.name, exc)
    source = {FormatType.TYPED: "text", FormatType.SCANNED: "tesseract", FormatType.HANDWRITTEN: "trocr"}[format_type]
    pages = []
    for i, (text, raw) in enumerate(extracted):
        width = height = url = None
        if i < len(image_paths):
            width, height = svc._image_size(image_paths[i])
            url = image_url(i) if image_url else None
        page = _page(i, text, url, width, height)
        calibrated = confidence_capture.calibrate(raw, format_type)
        page["lines"] = _text_lines(i, text, raw, calibrated, source)
        page["markdown"] = text
        pages.append(page)
    return Extraction(format_type, pages, page_count, source)


# ---------------------------------------------------------------------------------------------


def extract(
    path: Path, *, pages_dir: Path, format_hint: FormatType | None = None, pipeline: str = "ocr",
    lang: str = "en", image_url: ImageUrlFn | None = None, progress: ProgressFn | None = None,
    preprocess: dict[str, bool] | None = None,
) -> Extraction:
    """Read one file into contract pages; page images go to pages_dir/page-<n>.png.

    Raises ExtractionError (unreadable file), svc.PaddleOCRUnavailableError (the pipeline cannot
    run here), and the classic engines' OCRUnavailableError / HTRUnavailableError.
    """
    from pdfplumber.utils.exceptions import PdfminerException

    from app.modules import format_detection
    from app.modules.file_types import UnreadableFileError

    suffix = path.suffix.lower()
    if suffix in TYPED_SUFFIXES:
        _report(progress, 0.1, "Reading the text")
        return _typed_office(path)
    kind = "pdf" if suffix == ".pdf" else "image" if kind_of(path) is FileKind.IMAGE else None
    if kind is None:
        raise ExtractionError(f"{suffix} files are not supported")
    try:
        typed = kind == "pdf" and format_hint in (None, FormatType.TYPED) and format_detection.has_text_layer(str(path))
        if typed:
            return _typed_pdf(path, pages_dir, image_url, progress)
        if engine_mode() == "classic":
            return _classic_extract(path, kind, format_hint, pages_dir, image_url, progress)
    except (PdfminerException, UnreadableFileError) as exc:
        raise ExtractionError(f"not a readable {suffix} file: {exc}") from exc
    try:
        return _paddle_extract(
            path, kind, format_hint if format_hint is not FormatType.TYPED else None, pipeline, lang, pages_dir,
            image_url, progress, preprocess,
        )
    except svc.UnreadableInputError as exc:
        raise ExtractionError(str(exc)) from exc
