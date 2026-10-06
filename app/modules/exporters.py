"""Supports Module 7 — download an Experience Center result as TXT, JSON, Markdown, PDF, DOCX or CSV.

Responsibility: turn one result object (docs/experience_center_contract.md §4),
including the user's review edits, into a file:
- txt: recognised text, one line per OCR line, pages separated by a header
- layout (.layout.txt): the same text placed on a monospace grid from its line and word boxes,
  so indentation, column gaps (tables) and paragraph spacing look like the page
- json: the full result object
- md: the parsed Markdown (PP-StructureV3 / Office) or the lines as paragraphs
- pdf: searchable PDF: each page image with an invisible, selectable text layer
  placed exactly over the recognised lines (like a scanner's "OCR to PDF")
- docx: Word document built from the Markdown (headings, paragraphs, lists, tables)
- csv: the detected tables (one CSV, or a .zip of CSVs when there are several)

Uses from schemas.py: nothing; works on result dicts and page image paths.
"""

import csv
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

FORMATS = {
    "txt": ("text/plain; charset=utf-8", "txt"),
    "json": ("application/json", "json"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "layout": ("text/plain; charset=utf-8", "layout.txt"),
}

# Layout text: at most this many blank lines between two text rows (a big figure is not 40 empty lines).
MAX_BLANK_ROWS = 3

# Page images are rendered at paddle_pdf_dpi; this maps their pixels to PDF points.
POINTS_PER_INCH = 72


class ExportError(ValueError):
    """The result has nothing to export in the requested format."""


def page_text(page: dict) -> str:
    return "\n".join(line["text"] for line in page.get("lines") or [])


def result_markdown(result: dict) -> str:
    """The parsed Markdown, or (text OCR) the lines of each page as paragraphs."""
    pages = result.get("pages") or []
    if result.get("pipeline") != "ocr":
        parts = [page.get("markdown") or "" for page in pages]
        if any(parts):
            return "\n\n".join(p.strip() for p in parts if p.strip()) + "\n"
        if result.get("markdown"):
            return result["markdown"]
    return "\n\n".join(page_text(p) for p in pages if page_text(p)) + "\n"


def to_txt(result: dict) -> bytes:
    pages = result.get("pages") or []
    if len(pages) == 1:
        return (page_text(pages[0]) or result_markdown(result)).encode("utf-8")
    blocks = [f"===== Page {i + 1} =====\n{page_text(p) or (p.get('markdown') or '')}" for i, p in enumerate(pages)]
    return "\n\n".join(blocks).encode("utf-8")


def _median(values: list[float], default: float) -> float:
    values = sorted(v for v in values if v > 0)
    return values[len(values) // 2] if values else default


def _layout_items(line: dict) -> list[tuple[float, float, str]]:
    """(x0, x1, text) pieces of one line: its words when they still spell the line, else the whole line."""
    words = [w for w in line.get("words") or [] if w.get("box") and w.get("text")]
    if words and "".join(w["text"] for w in words) == "".join(line["text"].split()):
        return [(w["box"][0], w["box"][2], w["text"]) for w in words]
    return [(line["box"][0], line["box"][2], line["text"])]


# A gap wider than this many of the line's own characters is a column gap (kept); smaller is a space.
COLUMN_GAP_CHARS = 2.5


def page_layout_text(page: dict) -> str:
    """One page as monospace text: each line at its row, indentation and column gaps kept.

    Columns come from the page's median character width, rows from its median line height;
    gaps between rows become blank lines (up to MAX_BLANK_ROWS). Inside a line, words are
    joined by one space unless the gap is a real column gap (tables, tab stops), which is kept.
    """
    lines = [l for l in page.get("lines") or [] if l.get("box") and l.get("text", "").strip()]
    if not lines:
        return page.get("text") or page_text(page) or (page.get("markdown") or "")
    char_w = _median([(l["box"][2] - l["box"][0]) / max(1, len(l["text"])) for l in lines], 8.0)
    line_h = _median([l["box"][3] - l["box"][1] for l in lines], 12.0)
    left = min(l["box"][0] for l in lines)

    rows: list[list[dict]] = []  # lines that share a row (vertical overlap), top to bottom
    for line in sorted(lines, key=lambda l: (l["box"][1] + l["box"][3]) / 2):
        y0, y1 = line["box"][1], line["box"][3]
        if rows:
            top, bottom = min(l["box"][1] for l in rows[-1]), max(l["box"][3] for l in rows[-1])
            if min(bottom, y1) - max(top, y0) > 0.5 * min(bottom - top, y1 - y0):
                rows[-1].append(line)
                continue
        rows.append([line])

    out: list[str] = []
    previous_bottom = None
    for row in rows:
        top = min(l["box"][1] for l in row)
        if previous_bottom is not None:
            blanks = int(round((top - previous_bottom) / line_h))
            out.extend([""] * max(0, min(MAX_BLANK_ROWS, blanks)))
        previous_bottom = max(l["box"][3] for l in row)
        pieces = []
        for line in row:
            own_w = (line["box"][2] - line["box"][0]) / max(1, len(line["text"]))
            pieces.extend((x0, x1, text, own_w) for x0, x1, text in _layout_items(line))
        text, prev_x1 = "", None
        for x0, x1, piece, own_w in sorted(pieces, key=lambda i: i[0]):
            column = int(round((x0 - left) / char_w))
            if not text:
                text = " " * column + piece
            elif x0 - prev_x1 > COLUMN_GAP_CHARS * own_w:
                text = text.ljust(max(column, len(text) + 2)) + piece  # column gap: keep the position
            else:
                text += " " + piece
            prev_x1 = x1
        out.append(text.rstrip())
    return "\n".join(out)


def to_layout_txt(result: dict) -> bytes:
    pages = result.get("pages") or []
    if len(pages) == 1:
        return page_layout_text(pages[0]).encode("utf-8")
    blocks = [f"===== Page {i + 1} =====\n{page_layout_text(p)}" for i, p in enumerate(pages)]
    return "\n\n".join(blocks).encode("utf-8")


def to_json(result: dict) -> bytes:
    return json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")


def to_md(result: dict) -> bytes:
    return result_markdown(result).encode("utf-8")


# --- tables -------------------------------------------------------------------------------


class _TableParser(HTMLParser):
    """Collect the rows of every <table> in an HTML fragment (cells as plain text)."""

    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def html_tables(html: str) -> list[list[list[str]]]:
    parser = _TableParser()
    parser.feed(html)
    return [t for t in parser.tables if t]


def all_tables(result: dict) -> list[tuple[int, list[list[str]]]]:
    """(page index, rows) for every table; rows from the result, else parsed from its HTML."""
    found = []
    for i, page in enumerate(result.get("pages") or []):
        for table in page.get("tables") or []:
            rows = table.get("rows") or (html_tables(table.get("html") or "") or [[]])[0]
            if rows:
                found.append((i, rows))
    return found


def _csv_bytes(rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    csv.writer(buffer).writerows(rows)
    # BOM so Excel opens UTF-8 (accents, symbols) correctly.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def to_csv(result: dict) -> tuple[bytes, str, str]:
    """(content, media type, extension): one CSV, or a zip of table-<page>-<n>.csv files."""
    tables = all_tables(result)
    if not tables:
        raise ExportError("No tables were detected in this result (use Document parsing for tables).")
    if len(tables) == 1:
        return _csv_bytes(tables[0][1]), FORMATS["csv"][0], "csv"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        counts: dict[int, int] = {}
        for page_index, rows in tables:
            counts[page_index] = counts.get(page_index, 0) + 1
            archive.writestr(f"table-p{page_index + 1}-{counts[page_index]}.csv", _csv_bytes(rows))
    return buffer.getvalue(), "application/zip", "zip"


# --- searchable PDF -----------------------------------------------------------------------


def to_searchable_pdf(result: dict, page_images: list[Path | None], dpi: int) -> bytes:
    """Page images with an invisible text layer over each recognised line.

    Each line is written in render mode 3 (invisible) at its box, horizontally
    stretched to the box width, so selecting or searching text in a PDF viewer
    highlights the right spot on the scan.
    """
    import pymupdf

    pages = result.get("pages") or []
    if not pages or not any(page_images):
        raise ExportError("This result has no page images (Office files): export Markdown or DOCX instead.")
    scale = POINTS_PER_INCH / dpi
    doc = pymupdf.open()
    try:
        for page, image_path in zip(pages, page_images):
            if image_path is None or not Path(image_path).exists():
                continue
            width, height = page.get("width"), page.get("height")
            if not width or not height:
                with pymupdf.open(image_path) as img:
                    width, height = img[0].rect.width, img[0].rect.height
            pdf_page = doc.new_page(width=width * scale, height=height * scale)
            pdf_page.insert_image(pdf_page.rect, filename=str(image_path))
            for line in page.get("lines") or []:
                _insert_invisible_line(pdf_page, line, scale, pymupdf)
        if doc.page_count == 0:
            raise ExportError("The page images of this result are missing.")
        return doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()


def _insert_invisible_line(pdf_page, line: dict, scale: float, pymupdf) -> None:
    text = (line.get("text") or "").strip()
    box = line.get("box")
    if not text or not box:
        return
    x0, y0, x1, y1 = (v * scale for v in box)
    box_w, box_h = x1 - x0, y1 - y0
    if box_w <= 0 or box_h <= 0:
        return
    fontsize = max(1.0, box_h * 0.8)
    text_w = pymupdf.get_text_length(text, fontname="helv", fontsize=fontsize)
    origin = pymupdf.Point(x0, y1 - box_h * 0.2)  # baseline a little above the box bottom
    stretch = pymupdf.Matrix(box_w / text_w, 1) if text_w > 0 else pymupdf.Matrix(1, 1)
    # helv covers Latin-1 only; other characters become "?" in the hidden layer (still searchable for Latin text).
    safe = text.encode("latin-1", "replace").decode("latin-1")
    pdf_page.insert_text(origin, safe, fontsize=fontsize, fontname="helv", render_mode=3, morph=(origin, stretch))


# --- DOCX ---------------------------------------------------------------------------------

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_PIPE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_PIPE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_TAG = re.compile(r"<[^>]+>")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_inline(paragraph, text: str) -> None:
    """Add text to a python-docx paragraph, honouring **bold**, *italic* and `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_table(document, rows: list[list[str]]) -> None:
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
            if r == 0:
                for run in table.cell(r, c).paragraphs[0].runs:
                    run.bold = True


def _split_blocks(markdown: str) -> list[tuple[str, object]]:
    """Markdown -> [(kind, payload)]: heading, bullet, number, table, para. HTML tables become table blocks."""
    blocks: list[tuple[str, object]] = []
    # PP-StructureV3 emits tables (and centring divs) as HTML inside the Markdown.
    pieces = re.split(r"(<table.*?</table>)", markdown, flags=re.S | re.I)
    for piece in pieces:
        if piece.lower().startswith("<table"):
            for rows in html_tables(piece):
                blocks.append(("table", rows))
            continue
        lines = _TAG.sub("", piece).splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip():
                i += 1
                continue
            if _PIPE_ROW.match(line):
                rows = []
                while i < len(lines) and _PIPE_ROW.match(lines[i]):
                    if not _PIPE_RULE.match(lines[i]):
                        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                    i += 1
                blocks.append(("table", rows))
                continue
            if m := _HEADING.match(line):
                blocks.append(("heading", (len(m.group(1)), m.group(2).strip())))
            elif m := _BULLET.match(line):
                blocks.append(("bullet", m.group(1)))
            elif m := _NUMBERED.match(line):
                blocks.append(("number", m.group(1)))
            elif line.strip().startswith("![") or line.strip() == "---":
                pass  # images are not embedded; horizontal rules carry no text
            else:
                # Join consecutive plain lines into one paragraph.
                para = [line.strip()]
                while i + 1 < len(lines) and lines[i + 1].strip() and not any(
                    p.match(lines[i + 1]) for p in (_HEADING, _BULLET, _NUMBERED, _PIPE_ROW)
                ):
                    i += 1
                    para.append(lines[i].strip())
                blocks.append(("para", " ".join(para)))
            i += 1
    return blocks


def markdown_to_docx(markdown: str, title: str | None = None) -> bytes:
    from docx import Document

    document = Document()
    if title:
        document.core_properties.title = title
    for kind, payload in _split_blocks(markdown):
        if kind == "heading":
            level, text = payload
            document.add_heading(text, level=min(level, 4))
        elif kind == "bullet":
            _add_inline(document.add_paragraph(style="List Bullet"), payload)
        elif kind == "number":
            _add_inline(document.add_paragraph(style="List Number"), payload)
        elif kind == "table" and payload:
            _add_table(document, payload)
        elif kind == "para":
            _add_inline(document.add_paragraph(), payload)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_docx(result: dict) -> bytes:
    return markdown_to_docx(result_markdown(result), title=result.get("filename"))


def export(result: dict, fmt: str, page_images: list[Path | None] | None = None, dpi: int = 150) -> tuple[bytes, str, str]:
    """(content, media type, file extension) for one of FORMATS."""
    if fmt not in FORMATS:
        raise ExportError(f"Unknown export format {fmt!r}; use one of {', '.join(FORMATS)}")
    if fmt == "csv":
        return to_csv(result)
    media_type, ext = FORMATS[fmt]
    if fmt == "pdf":
        return to_searchable_pdf(result, page_images or [], dpi), media_type, ext
    builders = {"txt": to_txt, "json": to_json, "md": to_md, "docx": to_docx, "layout": to_layout_txt}
    return builders[fmt](result), media_type, ext
